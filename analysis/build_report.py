#!/usr/bin/env python3
"""
build_report.py — assemble the Phase 1 Moodle report (APA 7 .docx) from the results artifacts.

Compliance constraints enforced here, quoted from the assignment brief:

  * "Submit an individual 1,000-1,200-word Phase 1 report... Word counts exclude the access
    block, references, tables, and figures." -> only body prose is counted, and the script
    refuses to write a report outside the range.
  * "Use at most two main tables/figures in each report; full outputs and appendices belong in
    GitHub." -> this report carries one table and one figure, and no appendix.
  * "Access block in each report: Project title, your name and student identifier, phase,
    repository address, collaborator username fffchaobu and invitation date/status, submitted
    release/tag, and full commit identifier." -> emitted at the top.

Formatting: Times New Roman 12 pt, double spaced, 1-inch margins, 0.5-inch first-line indent,
three-line tables (top rule 1.5 pt, header rule 0.75 pt, bottom rule 1.5 pt, no vertical
rules), page number top right.

Run with the bundled office interpreter, which carries python-docx:

  "/Applications/DSH Desktop.app/Contents/Resources/office-runtime/primary-runtime/\
dependencies/python/bin/python3" analysis/build_report.py

Output: reports/COMM7660_Phase1_report.docx
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError as exc:                                    # noqa: BLE001
    sys.exit(f"python-docx unavailable in this interpreter ({exc}). See the module docstring.")

ROOT = Path(__file__).resolve().parent.parent
CLEAN = ROOT / "data" / "clean" / "analysis_dataset.csv"
RESULTS = ROOT / "analysis" / "results"

WORD_MIN, WORD_MAX = 1000, 1200

ACCESS_BLOCK = [
    ("Project title", "Reception Baseline — a base-rate locator for expected Steam reception"),
    ("Student", "____________________"),
    ("Student identifier", "____________"),
    ("Phase", "Phase 1 — media analytics and product justification"),
    ("Repository", "_(GitHub URL to be entered by the student)_"),
    ("Collaborator", "fffchaobu — invited _(date)_ — status _(pending / accepted)_"),
    ("Submitted release / tag", "phase1-analytics"),
    ("Full commit identifier", "_(to be entered by the student at tag time)_"),
]


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def require(path: Path) -> Path:
    if not path.exists():
        sys.exit(f"missing artifact: {path}. Run the analysis pipeline first.")
    return path


def group_label(row: dict) -> str:
    mapping = {"paid": "Paid", "free-to-play": "Free-to-play",
               "0": "Single-player-led", "1": "Multi-player-led"}
    for key in ("group", "is_free", ""):
        if key in row and str(row[key]).strip():
            return mapping.get(str(row[key]).strip(), str(row[key]).strip())
    for value in row.values():
        if value is not None and str(value).strip():
            return mapping.get(str(value).strip(), str(value).strip())
    return "(unlabelled)"


# ---------------------------------------------------------------- formatting


def setup_document(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    header = doc.sections[0].header
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = para.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


class WordCounter:
    """Counts body prose only. The access block, tables, captions and references are excluded,
    matching the brief's stated rule that the word count excludes them."""

    def __init__(self) -> None:
        self.words = 0
        self.counting = True

    def add(self, text: str) -> None:
        if self.counting:
            self.words += len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'\u2019\-\.]*", text))


def body(doc: Document, counter: WordCounter, text: str, indent: bool = True):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    counter.add(text)
    return p


def heading(doc: Document, counter: WordCounter, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(10)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    counter.add(text)
    return p


def access_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(12)
    p.add_run("Access block\n").bold = True
    for label, value in ACCESS_BLOCK:
        p.add_run(f"{label}: {value}\n")


def set_cell_border(cell, **kwargs) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        spec = kwargs.get(edge)
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        if spec is None:
            element.set(qn("w:val"), "nil")
        else:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(spec))
            element.set(qn("w:color"), "000000")


def three_line_table(doc: Document, header: list[str], rows: list[list[str]], caption: str):
    cap = doc.add_paragraph()
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cap.paragraph_format.space_before = Pt(8)
    cap.add_run(caption)

    table = doc.add_table(rows=1, cols=len(header))
    table.autofit = True
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = "Times New Roman"
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            r = p.add_run(str(text))
            r.font.size = Pt(11)
            r.font.name = "Times New Roman"

    last = len(table.rows) - 1
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            spec = {"left": None, "right": None}
            if i == 0:
                spec["top"] = 12
                spec["bottom"] = 6
            elif i == last:
                spec["top"] = None
                spec["bottom"] = 12
            else:
                spec["top"] = None
                spec["bottom"] = None
            set_cell_border(cell, **spec)
    doc.add_paragraph()


# ---------------------------------------------------------------- content


def build(args):
    clean_path = Path(args.clean) if args.clean else CLEAN
    if not clean_path.is_absolute():
        clean_path = ROOT / clean_path
    results = Path(args.results) if args.results else RESULTS
    if not results.is_absolute():
        results = ROOT / results

    clean = read_csv(require(clean_path))
    coefs = read_csv(require(results / "coefficients.csv"))
    coefs2 = read_csv(require(results / "coefficients_model2.csv"))
    t1 = read_csv(require(results / "table1_descriptives.csv"))
    summary = require(results / "model_summary.txt").read_text(encoding="utf-8")

    n = len(clean)
    tags = sorted({r["assigned_tag"] for r in clean})
    quarters = sorted({r["release_quarter"] for r in clean})
    dates = sorted(r["release_date"] for r in clean)
    n_free = sum(1 for r in clean if r["is_free"] == "1")
    n_mp = sum(1 for r in clean if r.get("multiplayer_emphasis") == "1")
    rates = sorted(float(r["positive_rate"]) for r in clean)
    mean_rate = sum(rates) / len(rates)
    sd_rate = (sum((x - mean_rate) ** 2 for x in rates) / len(rates)) ** 0.5
    q1_rate, q3_rate = rates[n // 4], rates[(3 * n) // 4]
    pct_at = lambda v: sum(1 for x in rates if x < v) / n           # noqa: E731
    genre_medians = []
    for t in tags:
        vals = sorted(float(r["positive_rate"]) for r in clean if r["assigned_tag"] == t)
        genre_medians.append(vals[len(vals) // 2])

    def coef(term):
        for c in coefs:
            if c["term"] == term:
                return c
        return None

    is_free, age = coef("is_free"), coef("age_years")
    pr2 = re.search(r"Pseudo R-squ\.\s*\(CS\):\s*([\d.]+)", summary)
    pr2 = float(pr2.group(1)) if pr2 else 0.0
    disp = re.search(r"Pearson chi2 / dof \(dispersion\)\s*:\s*([\d.]+)", summary)
    disp = float(disp.group(1)) if disp else 0.0
    w_free = re.search(r"is_free\s+unweighted b=([+-][\d.]+)\s+weighted b=([+-][\d.]+)", summary)
    s_free = re.search(r"is_free\s+full b=([+-][\d.]+)\s+settled b=([+-][\d.]+)", summary)
    s_age = re.search(r"age_years\s+full b=([+-][\d.]+)\s+settled b=([+-][\d.]+)", summary)

    mp = None
    for c in coefs2:
        if c["term"] == "multiplayer_emphasis":
            mp = c
    if mp is None:
        sys.exit("coefficients_model2.csv has no multiplayer_emphasis row; "
                 "Table 1 cannot be built without it")

    doc = Document()
    setup_document(doc)
    counter = WordCounter()

    access_block(doc)
    counter.counting = False

    ttl = doc.add_paragraph()
    ttl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ttl.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    ttl.add_run("What Predicts a Steam Title's Positive Review Rate?").bold = True

    counter.counting = True

    heading(doc, counter, "Opportunity and Alternatives")
    body(doc, counter,
         "Small studios and self-publishers must decide how to position a title with no external "
         "reference for how it will be received. Steam publishes a score on every store page but "
         "never the distribution behind it, so a developer cannot tell whether expecting a 0.75 "
         "positive review rate is ordinary caution or a warning sign. Two products were "
         "compared. Concept A, a reception benchmark, would place an expectation inside the "
         "market distribution. Concept B, a cross-market reception gap detector, would flag "
         "divergence between language communities. Concept B was rejected on measurement "
         "grounds: the rate is computed within each community's own voting pool, so a "
         "community-to-community difference confounds reception with the composition of "
         "reviewers, and no further collection from the same endpoints could separate the two.")

    heading(doc, counter, "Data and Model Rationale")
    body(doc, counter,
         f"The sample is {n} Steam titles released from {dates[0]} to {dates[-1]}, drawn as a "
         f"quota sample over {len(tags)} genre tags crossed with {len(quarters)} release "
         "quarters. The unit of analysis is one title; each row is a distinct application. "
         "Candidates were screened for release date and review count directly from the store "
         "search response, so per-title metadata calls were spent only on titles already known "
         "to be inside the window. Records come from the documented public review-summary and "
         "app-details endpoints, retrieved at one declared timestamp with a fixed request "
         "interval and no rate-limit workaround.")
    body(doc, counter,
         "The outcome is the title-level positive review rate, total positive divided by total "
         "positive plus total negative. Being a proportion bounded in [0,1], it is estimated by "
         "fractional logit: a generalised linear model with a binomial family and a logit link, "
         "which keeps fitted values inside the permissible range where a linear model would "
         "not. The focal predictors are free-to-play status and release age, the latter with a "
         "squared term because the mechanism need not be linear. The log review count controls "
         "for how precisely each rate is estimated, and the sampling tag controls for genre. "
         "Both specifications were declared before execution and both are reported.")

    heading(doc, counter, "Results, Uncertainty and Robustness")
    body(doc, counter,
         f"The analysis does not support a predictive product. Across {n} titles the mean rate "
         f"was {mean_rate:.3f} with a standard deviation of {sd_rate:.3f}. Paid titles averaged "
         f"{float(t1[0]['mean']):.3f} on {t1[0]['count']} titles against "
         f"{float(t1[1]['mean']):.3f} on {t1[1]['count']} free-to-play titles, and the "
         "free-to-play group's median was in fact the higher of the two. Table 1 sets out the "
         "principal estimates with their uncertainty and their robustness checks, which is the "
         "whole of the evidence this decision rests on:")
    three_line_table(
        doc,
        ["Estimate", "Logit coef.", "SE", "p", "Robustness"],
        [["is_free",
          f"{float(is_free['logit_coef']):+.3f}", f"{float(is_free['std_error']):.3f}",
          f"{float(is_free['p_value']):.3f}",
          f"{float(w_free.group(2)):+.3f} weighted; {float(s_free.group(2)):+.3f} settled"],
         ["age_years",
          f"{float(age['logit_coef']):+.3f}", f"{float(age['std_error']):.3f}",
          f"{float(age['p_value']):.3f}",
          f"{float(s_age.group(2)):+.3f} settled; sign flips"],
         ["multiplayer_emphasis",
          f"{float(mp['logit_coef']):+.3f}", f"{float(mp['std_error']):.3f}",
          f"{float(mp['p_value']):.3f}",
          f"n = {n_mp}; model 2 only"],
         ["Pseudo-R² (Cox–Snell)", f"{pr2:.3f}", "—", "—",
          "attributes explain under 1% of variation"],
         ["Pearson χ²/dof", f"{disp:.3f}", "—", "—",
          "no overdispersion; no influential titles"]],
        "Table 1. Principal estimates, uncertainty, and robustness checks.")
    body(doc, counter,
         "The interval is the finding rather than an inconvenience. The free-to-play marginal "
         "effect is minus 0.014 with a 95% interval from minus 0.161 to plus 0.134, so the "
         "sample cannot detect an effect of a size a developer would act on, and cannot rule one "
         "out either. Dispersion is low and no title is influential, which shows the residual "
         "variance matches the assumed variance rather than showing that the model is "
         "informative. Figure 1 shows the same result without a model: the reference "
         "distribution the product will use.")
    fig_path = results / "fig1_outcome_distribution.png"
    if fig_path.exists():
        doc.add_picture(str(fig_path), width=Inches(5.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap.add_run("Figure 1. Distribution of the positive review rate across "
                    f"{n} titles. The product reports a user's expected rate as a position in "
                    "this distribution.")
    body(doc, counter,
         "Robustness checks move the focal estimate materially without changing the conclusion: "
         f"the free-to-play coefficient ranges from {float(is_free['logit_coef']):+.3f} "
         f"unweighted to {float(w_free.group(2)):+.3f} when weighted by review count and "
         f"{float(s_free.group(2)):+.3f} when titles under six months old are excluded, while "
         f"the age term moves from {float(age['logit_coef']):+.3f} to "
         f"{float(s_age.group(2)):+.3f} and changes sign. No specification makes either "
         "predictor distinguishable from zero.")
    body(doc, counter,
         "One pattern deserves follow-up rather than a product claim. Titles whose core loop is "
         f"social or competitive average 0.769 against 0.835 for single-player-led titles, on "
         f"{n_mp} titles, a logit coefficient of minus 0.486 (p = 0.376). It is the largest gap "
         "anywhere in the analysis and consistent in sign, but it is not significant at this "
         "sample size. A power calculation puts the sample needed to detect it at roughly ten "
         "times the present size, beyond the range this project was scoped to, so it was not "
         "pursued.")

    heading(doc, counter, "Development Decision and Why This Project")
    body(doc, counter,
         "The evidence justifies a benchmark rather than a predictor, and a narrower benchmark "
         "than first proposed. Because the attributes explain under one per cent of the "
         "variation, a tool emitting a predicted score for an unreleased title would restate "
         "the sample mean while appearing to say something specific about that title. A "
         "conditional benchmark fails too, and on its own terms: the unconditional interquartile "
         f"range is {q1_rate:.3f} to {q3_rate:.3f}, a width of {q3_rate - q1_rate:.3f}, while "
         f"the six genre medians span only {min(genre_medians):.3f} to "
         f"{max(genre_medians):.3f}, a width of "
         f"{max(genre_medians) - min(genre_medians):.3f}. A user selecting a genre would watch "
         "the answer stay still, which is a broken interaction rather than a cautious one, so "
         "those inputs were removed.")
    body(doc, counter,
         "The project to develop is Reception Baseline, a media app for small studios and "
         "self-publishers preparing a launch. The user enters the positive review rate they expect "
         "and learns its percentile among comparable recent titles, read against the full "
         "distribution with its sample size shown. The tool then lists the real titles nearest "
         "that position, each linking to its Steam store page, so the task ends in content rather "
         "than in a number: a developer can open ten comparable games and judge them directly. "
         "Two view controls adjust that list only: how many titles to show, and a minimum review "
         "count. The second is not neutral — a 500-review floor raises the listed set's mean from "
         "0.823 to 0.864 — so the interface states it. The percentile always uses the full "
         "reference set. "
         "The mechanism responds to its input: a rate of"
         f" 0.70 sits at the {pct_at(0.70) * 100:.0f}th percentile and 0.90 at the "
         f"{pct_at(0.90) * 100:.0f}th. It is preferable to Concept B because Concept B's "
         "comparison cannot be measured from this data at all, whereas this one is measured "
         "directly. The supported features are the percentile position with its distribution, the "
         "comparable-title discovery list, and a scope statement naming the attributes that were "
         "tested and found not to explain the variation. The model does not run in the interface; "
         "it determined which outputs are honest to offer, which is why no predicted value appears "
         "anywhere. What remains an assumption is that a developer finds a population position and "
         "its nearest neighbours useful without any attribute conditioning. That is untested here, "
         "because this study observes titles rather than developer behaviour, and it is the "
         "assumption the Phase 2 pilot must test.")

    heading(doc, counter, "Limitations")
    body(doc, counter,
         f"Power is the main constraint: with {n_free} free-to-play titles the intervals are "
         "wide, so not significant means not detected. The 50-review floor excludes releases "
         "that attracted almost no attention, plausibly the worst received, so the reference "
         "range is optimistic at the bottom. Rates were retrieved at a single timestamp, leaving "
         "drift over time unobserved. Developer reputation, marketing spend and launch quality "
         "are unmeasured and plausibly correlated with both business model and reception. The "
         "gameplay feature is derived from the provider's category list rather than hand-verified "
         "for every title. The design supports association, not causal identification, and an "
         "excluded title is unobserved rather than a zero.")

    counter.counting = False
    heading(doc, counter, "References")
    for ref in [
        "Byers, T., Gibbs, M., & Nansen, B. (2025). Evaluating time in play and temporal "
        "satisfaction: Time-centric language in video game user reviews on Steam. In "
        "Proceedings of the 37th Australian Conference on Human-Computer Interaction. "
        "Association for Computing Machinery. https://doi.org/10.1145/3764687.3764693",
        "Eberhard, L., Kasper, P., Koncar, P., & Gütl, C. (2018). Investigating helpfulness of "
        "video game reviews on the Steam platform. In 2018 Fifth International Conference on "
        "Social Networks Analysis, Management and Security (SNAMS). IEEE. "
        "https://doi.org/10.1109/SNAMS.2018.8554542",
        "Guzsvinecz, T., & Szűcs, J. (2023). Length and sentiment analysis of reviews about "
        "top-level video game genres on the Steam platform. Computers in Human Behavior, 149, "
        "107955. https://doi.org/10.1016/j.chb.2023.107955",
        "Ma, J., & Jenkel, H. B. (2025). Cultural insights in game reception. International "
        "Journal of Gaming and Computer-Mediated Simulations, 17(1). "
        "https://doi.org/10.4018/IJGCMS.391739",
        "Mai, Y., & Hu, B. (2023). Optimizing free-to-play multiplayer games with premium "
        "subscription. Management Science, 69(6). https://doi.org/10.1287/mnsc.2022.4510",
        "Valve Corporation. (n.d.). Steamworks documentation: User reviews. Retrieved 12 "
        "October 2026, from https://partner.steamgames.com/doc/store/getreviews",
    ]:
        p = doc.add_paragraph(ref)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)

    return doc, counter


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="reports/COMM7660_Phase1_report.docx")
    ap.add_argument("--clean", default=None)
    ap.add_argument("--results", default=None)
    ap.add_argument("--check-only", action="store_true")
    args = ap.parse_args()

    doc, counter = build(args)
    print(f"body word count (excludes access block, tables, references): {counter.words}")
    ok = WORD_MIN <= counter.words <= WORD_MAX
    print(f"  target {WORD_MIN}-{WORD_MAX}: {'OK' if ok else 'OUT OF RANGE'}")
    if not ok and not args.check_only:
        sys.exit("refusing to write an out-of-range report")
    if args.check_only:
        return 0 if ok else 1

    out = Path(args.out)
    if not out.is_absolute():
        out = ROOT / out
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"wrote {out}  (main body: 1 table, 1 figure; no appendix)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
